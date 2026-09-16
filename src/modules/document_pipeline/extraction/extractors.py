"""Per-format text extractors.

One `TextExtractor` per supported format, mirroring how `formats.py` keeps per-format validation
knowledge in one registry. The shared principle: **read what the format already states, don't
infer it.** Office formats carry their structure explicitly (a Word heading is a named style, a
slide title is a title placeholder, a worksheet is already a grid), so headings and tables come
out of the file rather than out of a model. PDF is the one format with no semantic structure at
all — just positioned glyphs — so it gets a small, clearly-bounded font-size heuristic and
pdfplumber's ruling-line table detection.

No OCR: this corpus is digitally authored, so every file has a real text layer. A file that turns
out not to (a scan someone pasted in) produces almost no text here, which the normalization
quality gate flags explicitly rather than letting near-empty content through silently.
"""

import io
import logging
from abc import ABC, abstractmethod
from collections import Counter

from src.modules.document_pipeline.extraction.models import (
    ExtractedContent,
    ExtractedTable,
    ExtractedUnit,
    ExtractionError,
    Heading,
)
from src.modules.document_pipeline.formats import (
    DOCX_MIME,
    PDF_MIME,
    PPTX_MIME,
    XLSX_MIME,
)

logger = logging.getLogger(__name__)

# A PDF line qualifies as a heading if it is meaningfully larger than body text and short enough
# to be a title rather than a sentence that happens to be set large.
_HEADING_SIZE_RATIO = 1.15
_HEADING_MAX_CHARS = 120


class TextExtractor(ABC):
    """Extracts readable content from one format's bytes."""

    @abstractmethod
    def extract(self, data: bytes) -> ExtractedContent:
        """Returns the file's text, tables and headings, or raises `ExtractionError`."""


def _clean_cell(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _table_has_content(rows: list[list[str]]) -> bool:
    return any(any(cell for cell in row) for row in rows)


class PdfTextExtractor(TextExtractor):
    """PDF via pdfplumber — text layer, ruling-line tables, font-size heading heuristic."""

    def extract(self, data: bytes) -> ExtractedContent:
        import pdfplumber

        units: list[ExtractedUnit] = []
        tables: list[ExtractedTable] = []
        headings: list[Heading] = []
        # (text, font size, page index) for every line, collected across the whole document so
        # the body size is measured document-wide rather than per page — a page that is entirely
        # a heading would otherwise make that heading look like body text.
        lines: list[tuple[str, float, int]] = []

        try:
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                for index, page in enumerate(pdf.pages, start=1):
                    units.append(
                        ExtractedUnit(index=index, label=f"Page {index}", text=page.extract_text() or "")
                    )

                    for raw_table in page.extract_tables():
                        rows = [[_clean_cell(cell) for cell in row] for row in raw_table]
                        if _table_has_content(rows):
                            tables.append(ExtractedTable(unit_index=index, rows=rows))

                    lines.extend(self._page_lines(page, index))
        except ExtractionError:
            raise
        except Exception as e:
            raise ExtractionError(f"Could not read PDF content: {e}") from e

        headings = self._infer_headings(lines)
        return ExtractedContent(units=units, headings=headings, tables=tables)

    @staticmethod
    def _page_lines(page, page_index: int) -> list[tuple[str, float, int]]:
        """Groups a page's words into lines, each tagged with its largest font size."""
        try:
            words = page.extract_words(extra_attrs=["size"])
        except Exception:
            # Heading inference is a nicety; losing it must never fail the extraction itself.
            return []

        by_top: dict[int, list[dict]] = {}
        for word in words:
            # Round the vertical position so words on the same visual line group together despite
            # sub-pixel differences.
            by_top.setdefault(round(float(word.get("top", 0))), []).append(word)

        lines = []
        for _, line_words in sorted(by_top.items()):
            text = " ".join(w.get("text", "") for w in line_words).strip()
            if not text:
                continue
            size = max(float(w.get("size", 0) or 0) for w in line_words)
            lines.append((text, size, page_index))
        return lines

    @staticmethod
    def _infer_headings(lines: list[tuple[str, float, int]]) -> list[Heading]:
        """Treats lines notably larger than the document's most common font size as headings.

        This is the one place the pipeline infers structure instead of reading it, and it is
        deliberately conservative: a false negative just means a heading is treated as body text,
        which costs nothing downstream beyond slightly coarser chunk boundaries.
        """
        if not lines:
            return []

        # Weighted by characters, not by line count: body text is whichever size carries the most
        # *text*, which is what "body" means. Counting lines instead ties on short documents — a
        # page with one heading and one paragraph has one line of each, and the tie can hand
        # "body size" to the heading, after which nothing is large enough to be a heading at all.
        size_weights: Counter[float] = Counter()
        for text, size, _ in lines:
            if size > 0:
                size_weights[round(size, 1)] += len(text)
        if not size_weights:
            return []
        body_size = size_weights.most_common(1)[0][0]

        # Distinct heading sizes, largest first, become levels 1, 2, 3...
        heading_sizes = sorted(
            {round(size, 1) for _, size, _ in lines if size >= body_size * _HEADING_SIZE_RATIO},
            reverse=True,
        )
        levels = {size: level for level, size in enumerate(heading_sizes, start=1)}

        headings = []
        for text, size, page_index in lines:
            level = levels.get(round(size, 1))
            if level is not None and len(text) <= _HEADING_MAX_CHARS:
                headings.append(Heading(text=text, level=level, unit_index=page_index))
        return headings


class DocxTextExtractor(TextExtractor):
    """DOCX via python-docx. Headings come from Word's own paragraph styles."""

    def extract(self, data: bytes) -> ExtractedContent:
        from docx import Document as DocxDocument

        try:
            document = DocxDocument(io.BytesIO(data))
            paragraphs = []
            headings = []
            for paragraph in document.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                paragraphs.append(text)
                level = self._heading_level(paragraph)
                if level is not None:
                    headings.append(Heading(text=text, level=level, unit_index=1))

            tables = []
            for table in document.tables:
                rows = [[_clean_cell(cell.text) for cell in row.cells] for row in table.rows]
                if _table_has_content(rows):
                    tables.append(ExtractedTable(unit_index=1, rows=rows))
        except Exception as e:
            raise ExtractionError(f"Could not read Word document content: {e}") from e

        # One unit: a .docx has no pages until it is rendered, so there is nothing honest to
        # divide it by (the same reasoning that makes page_count None for DOCX at validation).
        units = [ExtractedUnit(index=1, label="Document", text="\n".join(paragraphs))]
        return ExtractedContent(units=units, headings=headings, tables=tables)

    @staticmethod
    def _heading_level(paragraph) -> int | None:
        """Word names built-in heading styles 'Heading 1'...'Heading 9'. Title is treated as
        level 1 since that is what it means structurally."""
        style_name = (getattr(paragraph.style, "name", "") or "").strip()
        if style_name.lower() == "title":
            return 1
        if not style_name.lower().startswith("heading"):
            return None
        tail = style_name[len("heading"):].strip()
        return int(tail) if tail.isdigit() else 1


class PptxTextExtractor(TextExtractor):
    """PPTX via python-pptx. One unit per slide; the slide's title placeholder is its heading."""

    def extract(self, data: bytes) -> ExtractedContent:
        from pptx import Presentation

        units: list[ExtractedUnit] = []
        headings: list[Heading] = []
        tables: list[ExtractedTable] = []

        try:
            presentation = Presentation(io.BytesIO(data))
            for index, slide in enumerate(presentation.slides, start=1):
                title = self._slide_title(slide)
                if title:
                    headings.append(Heading(text=title, level=1, unit_index=index))

                texts = []
                for shape in slide.shapes:
                    if getattr(shape, "has_table", False):
                        rows = [
                            [_clean_cell(cell.text) for cell in row.cells]
                            for row in shape.table.rows
                        ]
                        if _table_has_content(rows):
                            tables.append(ExtractedTable(unit_index=index, rows=rows))
                    elif getattr(shape, "has_text_frame", False):
                        text = shape.text_frame.text.strip()
                        if text:
                            texts.append(text)

                units.append(
                    ExtractedUnit(index=index, label=f"Slide {index}", text="\n".join(texts))
                )
        except Exception as e:
            raise ExtractionError(f"Could not read PowerPoint content: {e}") from e

        return ExtractedContent(units=units, headings=headings, tables=tables)

    @staticmethod
    def _slide_title(slide) -> str:
        try:
            title_shape = slide.shapes.title
        except Exception:
            return ""
        if title_shape is None or not getattr(title_shape, "has_text_frame", False):
            return ""
        return title_shape.text_frame.text.strip()


class XlsxTextExtractor(TextExtractor):
    """XLSX via openpyxl. A worksheet is already a grid, so each sheet is both a unit and a table.

    Opened with `data_only=True`: cached computed values are what a reader cares about, and it
    also means formula source (`=cmd|...`) never reaches the extracted text.
    """

    def extract(self, data: bytes) -> ExtractedContent:
        from openpyxl import load_workbook

        units: list[ExtractedUnit] = []
        headings: list[Heading] = []
        tables: list[ExtractedTable] = []

        workbook = None
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            for index, worksheet in enumerate(workbook.worksheets, start=1):
                rows = [
                    [_clean_cell(cell) for cell in row]
                    for row in worksheet.iter_rows(values_only=True)
                ]
                rows = [row for row in rows if any(row)]

                # The sheet name is the only structural label a spreadsheet carries.
                headings.append(Heading(text=worksheet.title, level=1, unit_index=index))
                if _table_has_content(rows):
                    tables.append(ExtractedTable(unit_index=index, rows=rows))

                text = "\n".join("\t".join(row).rstrip() for row in rows)
                units.append(
                    ExtractedUnit(index=index, label=f"Sheet: {worksheet.title}", text=text)
                )
        except Exception as e:
            raise ExtractionError(f"Could not read spreadsheet content: {e}") from e
        finally:
            if workbook is not None:
                try:
                    workbook.close()
                except Exception:
                    logger.debug("Workbook close failed; read-only handle will be collected")

        return ExtractedContent(units=units, headings=headings, tables=tables)


# Keyed by the same MIME constants validation uses, so a format cannot be accepted at upload and
# then have no way to be read here. `ExtractionService` asserts the two registries agree.
EXTRACTORS: dict[str, TextExtractor] = {
    PDF_MIME: PdfTextExtractor(),
    DOCX_MIME: DocxTextExtractor(),
    XLSX_MIME: XlsxTextExtractor(),
    PPTX_MIME: PptxTextExtractor(),
}

"""Extraction stage tests.

Fixtures are built in-process rather than committed as binaries — the same reasoning as
`test_formats.py`: what is under test stays readable in the diff, and a case like "a heading is
20pt and body text is 10pt" is expressed as data instead of an opaque file.
"""

import io
import uuid

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from src.modules.document_pipeline.extraction.models import (
    ExtractionError,
    ExtractionResult,
)
from src.modules.document_pipeline.extraction.service import (
    ExtractionService,
    unextractable_formats,
)
from src.modules.document_pipeline.formats import (
    DOCX_MIME,
    PDF_MIME,
    PPTX_MIME,
    XLSX_MIME,
)

DOC_ID = uuid.UUID("11111111-2222-3333-4444-555555555555")


# ==============================================================================
# Fixture builders
# ==============================================================================

def build_pdf(pages: list[list[tuple[str, int, int, int]]]) -> bytes:
    """Builds a minimal multi-page PDF. Each line is (text, font_size, x, y).

    Hand-built rather than pulled from a PDF-writing library so font sizes — the only signal the
    heading heuristic has — are set explicitly by the test rather than by a library's defaults.
    """
    page_count = len(pages)
    page_obj_ids = [3 + i * 2 for i in range(page_count)]
    content_obj_ids = [4 + i * 2 for i in range(page_count)]
    font_id = 3 + page_count * 2

    objects: list[tuple[int, bytes]] = [(1, b"<< /Type /Catalog /Pages 2 0 R >>")]
    kids = " ".join(f"{pid} 0 R" for pid in page_obj_ids).encode()
    objects.append(
        (2, b"<< /Type /Pages /Kids [" + kids + b"] /Count " + str(page_count).encode() + b" >>")
    )

    for i, lines in enumerate(pages):
        ops = []
        for text, size, x, y in lines:
            esc = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
            ops.append(f"BT /F1 {size} Tf {x} {y} Td ({esc}) Tj ET")
        content = "\n".join(ops).encode("latin-1")
        objects.append((page_obj_ids[i], (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_obj_ids[i]} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>"
        ).encode()))
        objects.append((
            content_obj_ids[i],
            b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        ))

    objects.append((font_id, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"))
    objects.sort(key=lambda o: o[0])

    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for obj_id, body in objects:
        offsets[obj_id] = len(out)
        out += f"{obj_id} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_pos = len(out)
    max_id = max(offsets)
    out += f"xref\n0 {max_id + 1}\n".encode() + b"0000000000 65535 f \n"
    for oid in range(1, max_id + 1):
        out += f"{offsets.get(oid, 0):010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {max_id + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF".encode()
    return bytes(out)


def build_docx(with_table: bool = True) -> bytes:
    document = DocxDocument()
    document.add_heading("Air Quality Directive 2026", level=1)
    document.add_paragraph("This directive sets out obligations for stubble management.")
    document.add_heading("Section A: Scope", level=2)
    document.add_paragraph("Applies to all districts in the NCR region.")
    if with_table:
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "District"
        table.rows[0].cells[1].text = "Target"
        table.rows[1].cells[0].text = "Ludhiana"
        table.rows[1].cells[1].text = "40%"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Emissions"
    sheet.append(["Site", "PM2.5"])
    sheet.append(["Anand Vihar", 312])
    second = workbook.create_sheet("Targets")
    second.append(["Metric", "Goal"])
    second.append(["AQI", 150])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Quarterly Review"
    slide.placeholders[1].text = "Stubble burning down 22% YoY"

    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "Data"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Month"
    table.cell(0, 1).text = "Count"
    table.cell(1, 0).text = "Jan"
    table.cell(1, 1).text = "1204"

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def service() -> ExtractionService:
    return ExtractionService()


# ==============================================================================
# Every accepted format must be extractable
# ==============================================================================

def test_every_uploadable_format_has_an_extractor():
    """A format accepted at validation with no extractor would sit at EXTRACTION_FAILED forever,
    so adding one to formats.py without one here must fail loudly right here."""
    assert unextractable_formats() == set()


# ==============================================================================
# PDF
# ==============================================================================

def test_pdf_extracts_text_per_page(service):
    pdf = build_pdf([
        [("Policy overview", 10, 72, 720)],
        [("Second page body text", 10, 72, 720)],
    ])
    result = service.extract(DOC_ID, pdf, PDF_MIME)

    assert result.unit_count == 2
    assert [u.label for u in result.units] == ["Page 1", "Page 2"]
    assert "Policy overview" in result.units[0].text
    assert "Second page" in result.units[1].text
    assert result.char_count > 0


def test_pdf_infers_headings_from_relative_font_size(service):
    """The one place the pipeline infers structure rather than reading it — body text must not be
    swept up as a heading just because a document has few lines."""
    pdf = build_pdf([
        [
            ("Air Quality Directive 2026", 20, 72, 720),
            ("This directive sets out obligations for stubble management", 10, 72, 690),
            ("Section A Scope", 14, 72, 660),
            ("Applies to all districts in the NCR region", 10, 72, 630),
        ],
        [
            ("Section B Penalties", 14, 72, 720),
            ("Non compliance attracts a fine under section 15", 10, 72, 690),
        ],
    ])
    result = service.extract(DOC_ID, pdf, PDF_MIME)

    by_text = {h.text: h for h in result.headings}
    assert by_text["Air Quality Directive 2026"].level == 1
    assert by_text["Section A Scope"].level == 2
    # Body text is the most common size, so it is the baseline and never a heading.
    assert not any("obligations for stubble" in h.text for h in result.headings)
    # Body size is measured across the whole document, so a heading alone on a later page is
    # still recognised relative to the document's body text, not that page's.
    assert by_text["Section B Penalties"].unit_index == 2


def test_pdf_heading_survives_a_one_line_body(service):
    """Body size is weighted by characters, not line count. Counting lines ties here — one
    heading line, one body line — and the tie can hand 'body size' to the heading, after which
    nothing is large enough to be a heading at all."""
    pdf = build_pdf([[("Heading", 20, 72, 720), ("body text here", 10, 72, 690)]])
    result = service.extract(DOC_ID, pdf, PDF_MIME)

    assert [h.text for h in result.headings] == ["Heading"]


def test_pdf_with_no_text_layer_extracts_cleanly_but_empty(service):
    """A scan (or a blank page) is not an extraction *failure* — it reads fine, there is just
    nothing in it. Flagging that is the normalization quality gate's job, not this stage's."""
    pdf = build_pdf([[]])
    result = service.extract(DOC_ID, pdf, PDF_MIME)

    assert result.unit_count == 1
    assert result.char_count == 0
    assert result.headings == []


def test_unreadable_pdf_raises_extraction_error(service):
    with pytest.raises(ExtractionError):
        service.extract(DOC_ID, b"%PDF-1.4\nnot actually a pdf body", PDF_MIME)


# ==============================================================================
# DOCX
# ==============================================================================

def test_docx_reads_headings_from_word_styles(service):
    result = service.extract(DOC_ID, build_docx(), DOCX_MIME)

    levels = {h.text: h.level for h in result.headings}
    assert levels["Air Quality Directive 2026"] == 1
    assert levels["Section A: Scope"] == 2
    assert "stubble management" in result.full_text


def test_docx_is_a_single_unit(service):
    """Word text reflows — there are no pages until it is rendered, which is the same reason
    page_count is None for DOCX at validation."""
    result = service.extract(DOC_ID, build_docx(), DOCX_MIME)

    assert result.unit_count == 1
    assert result.units[0].label == "Document"


def test_docx_tables_are_kept_as_rows(service):
    result = service.extract(DOC_ID, build_docx(), DOCX_MIME)

    assert len(result.tables) == 1
    assert result.tables[0].rows == [["District", "Target"], ["Ludhiana", "40%"]]
    assert result.tables[0].row_count == 2


# ==============================================================================
# XLSX
# ==============================================================================

def test_xlsx_makes_one_unit_per_sheet(service):
    result = service.extract(DOC_ID, build_xlsx(), XLSX_MIME)

    assert result.unit_count == 2
    assert [u.label for u in result.units] == ["Sheet: Emissions", "Sheet: Targets"]
    assert [h.text for h in result.headings] == ["Emissions", "Targets"]


def test_xlsx_sheets_become_tables_with_values(service):
    result = service.extract(DOC_ID, build_xlsx(), XLSX_MIME)

    assert result.tables[0].rows == [["Site", "PM2.5"], ["Anand Vihar", "312"]]
    assert result.tables[1].rows == [["Metric", "Goal"], ["AQI", "150"]]


# ==============================================================================
# PPTX
# ==============================================================================

def test_pptx_makes_one_unit_per_slide_with_title_headings(service):
    result = service.extract(DOC_ID, build_pptx(), PPTX_MIME)

    assert result.unit_count == 2
    assert [u.label for u in result.units] == ["Slide 1", "Slide 2"]
    assert [h.text for h in result.headings] == ["Quarterly Review", "Data"]
    assert "22% YoY" in result.full_text


def test_pptx_tables_are_extracted(service):
    result = service.extract(DOC_ID, build_pptx(), PPTX_MIME)

    assert len(result.tables) == 1
    assert result.tables[0].rows == [["Month", "Count"], ["Jan", "1204"]]
    assert result.tables[0].unit_index == 2


# ==============================================================================
# Service behaviour and the stored artifact
# ==============================================================================

def test_unsupported_mime_raises_rather_than_guessing(service):
    with pytest.raises(ExtractionError):
        service.extract(DOC_ID, build_docx(), "image/png")


def test_result_survives_a_json_round_trip(service):
    """The result is stored as extracted/{id}.json and read back by the normalization stage, so
    a field that doesn't survive serialization would break the handoff silently."""
    result = service.extract(DOC_ID, build_docx(), DOCX_MIME)
    restored = ExtractionResult.model_validate_json(result.model_dump_json())

    assert restored.document_id == result.document_id
    assert restored.full_text == result.full_text
    assert restored.char_count == result.char_count
    assert [t.rows for t in restored.tables] == [t.rows for t in result.tables]


def test_full_text_is_derived_from_units_not_stored_separately(service):
    """full_text is computed, so it cannot drift out of sync with units."""
    result = service.extract(DOC_ID, build_xlsx(), XLSX_MIME)

    assert result.full_text == "\n\n".join(u.text for u in result.units if u.text)
